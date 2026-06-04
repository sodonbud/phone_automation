"""Excel read/write client using openpyxl.

Keeps a single Workbook open for the run's lifetime; caller must call .close()
(or use as a context manager) to guarantee the final save.
"""

from __future__ import annotations

import os
from datetime import datetime
from typing import Any

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from logger import get_logger

# Expected column names (order must match the sheet header row)
COLUMNS = [
    "Step",
    "Action",
    "Target Phone",
    "Number/Code",
    "Value",
    "Expected Result",
    "Actual Result",
    "Pass/Fail",
    "Notes",
    "Timestamp",
]

PASS_FILL = PatternFill("solid", fgColor="00B050")
FAIL_FILL = PatternFill("solid", fgColor="FF0000")
SKIP_FILL = PatternFill("solid", fgColor="FFFF00")
_RESULT_FILLS = {"PASS": PASS_FILL, "FAIL": FAIL_FILL, "SKIP": SKIP_FILL}

# Column indices (1-based) for result columns
_COL_ACTUAL = COLUMNS.index("Actual Result") + 1
_COL_PASS_FAIL = COLUMNS.index("Pass/Fail") + 1
_COL_NOTES = COLUMNS.index("Notes") + 1
_COL_TIMESTAMP = COLUMNS.index("Timestamp") + 1


class ExcelClient:
    def __init__(self, file_path: str, sheet_name: str) -> None:
        self._path = file_path
        self._sheet_name = sheet_name
        self._wb: openpyxl.Workbook | None = None
        self._ws = None
        self._header_map: dict[str, int] = {}  # column name → 1-based col index
        self.log = get_logger()

    # ------------------------------------------------------------------
    # Context manager support
    # ------------------------------------------------------------------

    def __enter__(self) -> "ExcelClient":
        self.open()
        return self

    def __exit__(self, *_: Any) -> None:
        self.close()

    # ------------------------------------------------------------------
    # Lifecycle
    # ------------------------------------------------------------------

    def open(self) -> None:
        if not os.path.exists(self._path):
            raise FileNotFoundError(f"Excel file not found: {self._path}")
        try:
            self._wb = openpyxl.load_workbook(
                self._path, keep_vba=False, data_only=True
            )
        except Exception as exc:
            raise RuntimeError(f"Cannot open Excel file '{self._path}': {exc}") from exc

        if self._sheet_name not in self._wb.sheetnames:
            available = ", ".join(self._wb.sheetnames)
            raise ValueError(
                f"Sheet '{self._sheet_name}' not found. Available sheets: {available}"
            )

        self._ws = self._wb[self._sheet_name]
        self._build_header_map()
        self.log.info("Opened '%s' sheet '%s'", self._path, self._sheet_name)

    def close(self) -> None:
        if self._wb is not None:
            try:
                self._wb.save(self._path)
                self.log.info("Workbook saved to '%s'", self._path)
            except Exception as exc:
                self.log.error("Failed to save workbook: %s", exc)
            finally:
                self._wb = None
                self._ws = None

    # ------------------------------------------------------------------
    # Reading
    # ------------------------------------------------------------------

    def _build_header_map(self) -> None:
        header_row = next(self._ws.iter_rows(min_row=1, max_row=1, values_only=True))
        self._header_map = {
            str(cell).strip(): idx + 1
            for idx, cell in enumerate(header_row)
            if cell is not None
        }

    def get_test_cases(self) -> list[dict[str, Any]]:
        """Return all data rows as a list of dicts keyed by header name."""
        rows = []
        for row_idx, row in enumerate(
            self._ws.iter_rows(min_row=2, values_only=True), start=2
        ):
            # Skip completely blank rows
            if all(cell is None or str(cell).strip() == "" for cell in row):
                continue
            record: dict[str, Any] = {"_row": row_idx}
            for name, col in self._header_map.items():
                val = row[col - 1] if col - 1 < len(row) else None
                record[name] = "" if val is None else str(val).strip()
            rows.append(record)
        self.log.info("Loaded %d test case(s)", len(rows))
        return rows

    # ------------------------------------------------------------------
    # Writing
    # ------------------------------------------------------------------

    def write_result(
        self,
        row_index: int,
        actual: str,
        result: str,
        notes: str,
        timestamp: str | None = None,
    ) -> None:
        """Update Actual Result, Pass/Fail, Notes, Timestamp for *row_index* (1-based)."""
        if self._ws is None:
            raise RuntimeError("Workbook is not open.")
        if timestamp is None:
            timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

        result_upper = result.strip().upper()
        fill = _RESULT_FILLS.get(result_upper)

        def _set(col: int, value: Any, apply_fill: bool = False) -> None:
            cell = self._ws.cell(row=row_index, column=col)
            cell.value = value
            if apply_fill and fill:
                cell.fill = fill

        _set(_COL_ACTUAL, actual)
        _set(_COL_PASS_FAIL, result_upper, apply_fill=True)
        _set(_COL_NOTES, notes)
        _set(_COL_TIMESTAMP, timestamp)

        # Persist immediately so a crash doesn't lose earlier results
        try:
            self._wb.save(self._path)
        except PermissionError:
            self.log.error(
                "Cannot save '%s' — close the file in Excel first, then re-run.", self._path
            )
        except Exception as exc:
            self.log.error("Failed to save after row %d: %s", row_index, exc)

    # ------------------------------------------------------------------
    # Template generation (used by generate_template.py / main.py)
    # ------------------------------------------------------------------

    @staticmethod
    def create_template(file_path: str, sheet_name: str = "TestCases") -> None:
        """Write a ready-to-use test sheet to *file_path*."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = sheet_name

        # ── Styles ────────────────────────────────────────────────────
        header_font  = Font(bold=True, color="FFFFFF")
        header_fill  = PatternFill("solid", fgColor="4472C4")
        header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

        section_font  = Font(bold=True, color="FFFFFF")
        section_fills = {
            "CALL": PatternFill("solid", fgColor="375623"),   # dark green
            "SMS" : PatternFill("solid", fgColor="1F4E79"),   # dark blue
            "USSD": PatternFill("solid", fgColor="7B2C2C"),   # dark red
        }

        row_fill_a = PatternFill("solid", fgColor="EBF3E8")   # light green
        row_fill_b = PatternFill("solid", fgColor="DDEEFF")   # light blue
        row_fill_u = PatternFill("solid", fgColor="FFF2CC")   # light yellow

        center = Alignment(horizontal="center", vertical="center")
        left   = Alignment(horizontal="left",   vertical="center", wrap_text=True)

        # ── Header row ────────────────────────────────────────────────
        for col_idx, name in enumerate(COLUMNS, start=1):
            cell = ws.cell(row=1, column=col_idx, value=name)
            cell.font      = header_font
            cell.fill      = header_fill
            cell.alignment = header_align
        ws.row_dimensions[1].height = 28

        # ── Column widths ─────────────────────────────────────────────
        widths = [5, 14, 12, 22, 24, 28, 28, 10, 32, 20]
        for col_idx, width in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(col_idx)].width = width

        # ── Data rows definition ──────────────────────────────────────
        # fmt: (step, action, target, number_code, value, expected, section)
        rows = [
            # ── CALL TEST ──────────────────────────────────────────────
            ("SECTION", "CALL TEST — Phone1 calls Phone2, Phone2 answers, both record"),
            ("1",  "CALL",         "Phone1", "+97699002222", "",                     "",                  "CALL"),
            ("2",  "ANSWER_CALL",  "Phone2", "",             "5",                    "Call answered",     "CALL"),
            ("3",  "START_RECORD", "Phone1", "",             "recordings/p1_call.wav","",                 "CALL"),
            ("4",  "START_RECORD", "Phone2", "",             "recordings/p2_call.wav","",                 "CALL"),
            ("5",  "WAIT",         "Phone1", "10",           "",                     "",                  "CALL"),
            ("6",  "END_CALL",     "Phone1", "",             "",                     "",                  "CALL"),
            ("7",  "STOP_RECORD",  "Phone1", "",             "recordings/p1_call.wav","",                 "CALL"),
            ("8",  "STOP_RECORD",  "Phone2", "",             "recordings/p2_call.wav","",                 "CALL"),
            # ── SMS TEST ───────────────────────────────────────────────
            ("SECTION", "SMS TEST — Phone1 sends SMS, verify it arrives on Phone2"),
            ("9",  "SMS",          "Phone1", "+97699002222", "Hello from Phone1",    "",                  "SMS"),
            ("10", "WAIT",         "Phone1", "15",           "",                     "",                  "SMS"),
            ("11", "CHECK_SMS",    "Phone2", "+97699001111", "Hello from Phone1",    "Hello from Phone1", "SMS"),
            # ── USSD TEST ──────────────────────────────────────────────
            ("SECTION", "USSD TEST — Dial USSD on Phone1, capture network response"),
            ("12", "USSD",         "Phone1", "*100#",        "",                     "",                  "USSD"),
        ]

        # ── Notes for each action ─────────────────────────────────────
        action_notes = {
            "CALL":         "Phone1 dials Phone2. Leave Expected Result blank.",
            "ANSWER_CALL":  "Waits up to Value(secs) for incoming call on Phone2, then answers automatically.",
            "START_RECORD": "Starts call recording. File saved to path in Value column.",
            "WAIT":         "Pause Number/Code seconds between steps.",
            "END_CALL":     "Hangs up the call on Phone1.",
            "STOP_RECORD":  "Stops recording and pulls audio file to local path in Value column.",
            "SMS":          "Sends SMS from Phone1 to Number/Code with body in Value.",
            "CHECK_SMS":    "Queries Phone2 inbox for SMS from Number/Code containing Value text.",
            "USSD":         "Dials USSD code and captures network response text.",
        }

        section_fills_map = {"CALL": row_fill_a, "SMS": row_fill_b, "USSD": row_fill_u}

        data_row = 2
        for entry in rows:
            if entry[0] == "SECTION":
                # Section header spanning all columns
                label   = entry[1]
                section = label.split()[0]
                ws.merge_cells(start_row=data_row, start_column=1,
                               end_row=data_row, end_column=len(COLUMNS))
                cell = ws.cell(row=data_row, column=1, value=f"  {label}")
                cell.font      = section_font
                cell.fill      = section_fills.get(section, header_fill)
                cell.alignment = Alignment(vertical="center")
                ws.row_dimensions[data_row].height = 22
                data_row += 1
                continue

            step, action, target, num_code, value, expected, section = entry
            row_fill = section_fills_map.get(section, row_fill_a)
            note     = action_notes.get(action, "")

            values = [step, action, target, num_code, value, expected, "", "", note, ""]
            for col_idx, val in enumerate(values, start=1):
                cell = ws.cell(row=data_row, column=col_idx, value=val)
                cell.fill      = row_fill
                cell.alignment = center if col_idx in (1, 2, 3, 8) else left
            ws.row_dimensions[data_row].height = 18
            data_row += 1

        # ── Freeze header, auto-filter ────────────────────────────────
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}1"

        wb.save(file_path)
        print(f"Template written to '{file_path}'")

